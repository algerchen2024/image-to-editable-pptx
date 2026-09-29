from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import crop_asset  # noqa: E402
import inspect_pptx  # noqa: E402
import ocr_to_pageir  # noqa: E402
import pageir_common  # noqa: E402
import sample_colors  # noqa: E402
import validate_page_ir  # noqa: E402


class CommonTests(unittest.TestCase):
    def test_normalize_hex(self):
        self.assertEqual(pageir_common.normalize_hex("#fa0"), "#FFAA00")
        self.assertEqual(pageir_common.normalize_hex("#12abEF"), "#12ABEF")
        with self.assertRaises(ValueError):
            pageir_common.normalize_hex("#12")

    def test_slide_size_landscape_and_portrait(self):
        w, h = pageir_common.slide_size_inches(1600, 900)
        self.assertAlmostEqual(w, 13.333333)
        self.assertAlmostEqual(w / h, 1600 / 900, places=6)
        w, h = pageir_common.slide_size_inches(900, 1600)
        self.assertAlmostEqual(h, 7.5)
        self.assertAlmostEqual(w / h, 900 / 1600, places=6)

    def test_pt_per_px(self):
        # 1600 px across 13.333 in -> 1 px = 0.6 pt.
        self.assertAlmostEqual(pageir_common.pt_per_px(1600, 900), 0.6, places=4)


class SampleColorTests(unittest.TestCase):
    def test_background_and_foreground(self):
        image = np.full((100, 200, 3), 255, dtype=np.uint8)
        image[40:60, 50:150] = (240, 90, 40)  # orange "ink" inside a white box
        sample = sample_colors.sample_box(image, [30, 20, 140, 60])
        self.assertEqual(sample["background"], "#FFFFFF")
        self.assertEqual(sample["foreground"], "#F05A28")
        self.assertEqual(sample_colors.page_background(image), "#FFFFFF")

    def test_bbox_outside_image(self):
        image = np.zeros((10, 10, 3), dtype=np.uint8)
        with self.assertRaises(ValueError):
            sample_colors.sample_box(image, [20, 20, 5, 5])


class OcrToPageIrTests(unittest.TestCase):
    def test_font_estimate_cjk(self):
        # CJK ink ~0.88 em; 44 px ink on a 1600 px wide page -> 50 px em -> 30 pt.
        self.assertEqual(ocr_to_pageir.estimate_font_pt("标题文字", 44, 1600, 900), 30.0)

    def test_draft_is_valid_pageir(self):
        ocr = {
            "width_px": 1600,
            "height_px": 900,
            "lines": [
                {"text": "标题文字", "bbox": [70, 60, 400, 44], "confidence": 92},
                {"text": "Quarterly plan", "bbox": [70, 200, 300, 30], "confidence": 50},
                {"text": "  ", "bbox": [0, 0, 1, 1], "confidence": 99},
            ],
        }
        image = np.full((900, 1600, 3), 255, dtype=np.uint8)
        image[60:104, 70:470] = (17, 17, 17)
        payload = ocr_to_pageir.build_page_ir(ocr, {"cjk": "Microsoft YaHei", "latin": "Arial"}, image)
        objects = payload["pages"][0]["objects"]
        self.assertEqual(len(objects), 2)
        self.assertEqual(objects[0]["style"]["font_face"], "Microsoft YaHei")
        self.assertEqual(objects[0]["style"]["color"], "#111111")
        self.assertIn("LOW", objects[1]["note"])
        self.assertEqual(payload["pages"][0]["background"], "#FFFFFF")
        self.assertEqual(validate_page_ir.validate_page_ir(payload, ROOT), [])


class CropAssetTests(unittest.TestCase):
    def test_crop_pad_clamp_and_alpha(self):
        image = Image.new("RGB", (100, 80), "white")
        image.paste((0, 0, 255), (20, 20, 40, 40))
        asset, bbox = crop_asset.crop_asset(image, [10, 10, 40, 40], pad=15, white_to_alpha=245)
        self.assertEqual(bbox, [0, 0, 65, 65])
        self.assertEqual(asset.size, (65, 65))
        self.assertEqual(asset.getpixel((0, 0))[3], 0)  # white became transparent
        self.assertEqual(asset.getpixel((25, 25)), (0, 0, 255, 255))


class InspectCoverageTests(unittest.TestCase):
    def test_tiled_pictures_are_detected(self):
        tiles = [(0, 0, 500, 500), (500, 0, 500, 500), (0, 500, 500, 500), (500, 500, 500, 500)]
        self.assertGreaterEqual(inspect_pptx.picture_union_coverage(tiles, 1000, 1000), 0.99)
        self.assertAlmostEqual(inspect_pptx.picture_union_coverage(tiles[:1], 1000, 1000), 0.25, places=2)

    def test_tiled_screenshot_fails_inspection(self):
        from pptx import Presentation
        from pptx.util import Emu

        with tempfile.TemporaryDirectory() as tmp:
            tile = Path(tmp) / "tile.png"
            Image.new("RGB", (10, 10), "gray").save(tile)
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            half_w, half_h = prs.slide_width // 2, prs.slide_height // 2
            for x in (0, half_w):
                for y in (0, half_h):
                    slide.shapes.add_picture(str(tile), Emu(x), Emu(y), Emu(half_w), Emu(half_h))
            out = Path(tmp) / "tiled.pptx"
            prs.save(out)
            report = inspect_pptx.inspect(out)
        self.assertFalse(report["pass"])
        self.assertEqual(report["full_slide_picture_count"], 0)
        self.assertEqual(report["raster_shortcut_slides"], [1])


if __name__ == "__main__":
    unittest.main()
