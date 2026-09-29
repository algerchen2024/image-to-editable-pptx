"""Compile PageIR with the Node compiler and read the result back with python-pptx."""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu


ROOT = Path(__file__).resolve().parents[1]
COMPILER = ROOT / "scripts" / "compile_page_ir.js"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
EMU_PER_IN = 914400


def node_ready() -> bool:
    node = shutil.which("node")
    if not node:
        return False
    proc = subprocess.run([node, "-e", "require('pptxgenjs')"], cwd=ROOT, capture_output=True, check=False)
    return proc.returncode == 0


def compile_ir(payload: dict, tmp: Path) -> Presentation:
    ir = tmp / "page_ir.json"
    ir.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    out = tmp / "out.pptx"
    proc = subprocess.run(["node", str(COMPILER), str(ir), str(out)], cwd=ROOT, capture_output=True, text=True, check=False)
    if proc.returncode != 0:
        raise AssertionError(proc.stderr or proc.stdout)
    return Presentation(out)


def shape_by_name_order(prs: Presentation):
    return list(prs.slides[0].shapes)


@unittest.skipUnless(node_ready(), "node with pptxgenjs is not available (run npm ci)")
class CompileRoundTripTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_minimal_example_compiles(self):
        payload = json.loads((ROOT / "examples" / "minimal" / "page_ir.json").read_text(encoding="utf-8"))
        prs = compile_ir(payload, self.tmp)
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 1600 / 900, places=3)

    def test_reversed_lines_keep_direction(self):
        payload = page(
            [
                {"id": "left", "type": "line", "points": [800, 100, 200, 100], "z": 1, "style": {"end_arrow": "triangle"}},
                {"id": "up", "type": "line", "points": [500, 800, 300, 400], "z": 2, "style": {"end_arrow": "triangle"}},
                {"id": "down", "type": "line", "points": [100, 100, 300, 500], "z": 3},
            ]
        )
        prs = compile_ir(payload, self.tmp)
        left, up, down = shape_by_name_order(prs)
        scale = prs.slide_width / 1600
        for shape in (left, up, down):
            self.assertGreaterEqual(shape.width, 0)
            self.assertGreaterEqual(shape.height, 0)
        self.assertAlmostEqual(left.left, 200 * scale, delta=EMU_PER_IN * 0.01)
        self.assertAlmostEqual(left.width, 600 * scale, delta=EMU_PER_IN * 0.01)
        self.assertEqual(xfrm(left).get("flipH"), "1")
        self.assertEqual(xfrm(up).get("flipH"), "1")
        self.assertEqual(xfrm(up).get("flipV"), "1")
        self.assertIsNone(xfrm(down).get("flipH"))
        self.assertIsNone(xfrm(down).get("flipV"))

    def test_short_hex_dash_and_corner_radius(self):
        payload = page(
            [
                {
                    "id": "card",
                    "type": "round_rect",
                    "bbox": [100, 100, 400, 200],
                    "style": {"fill": "#FA0", "line": "#123", "line_dash": "dash", "corner_radius_px": 50},
                },
                {"id": "dotted", "type": "line", "points": [100, 600, 700, 600], "style": {"dash": "dot"}},
            ]
        )
        prs = compile_ir(payload, self.tmp)
        card, dotted = shape_by_name_order(prs)
        xml = card._element
        self.assertEqual(xml.find(".//a:solidFill/a:srgbClr", NS).get("val"), "FFAA00")
        self.assertEqual(xml.find(".//a:ln//a:srgbClr", NS).get("val"), "112233")
        self.assertEqual(xml.find(".//a:ln/a:prstDash", NS).get("val"), "dash")
        # 50 px radius on a 200 px tall box -> adj = 50/200 = 25000.
        adj = xml.find(".//a:prstGeom/a:avLst/a:gd", NS)
        self.assertIsNotNone(adj)
        self.assertAlmostEqual(int(adj.get("fmla").split()[1]), 25000, delta=500)
        self.assertEqual(dotted._element.find(".//a:ln/a:prstDash", NS).get("val"), "sysDot")

    def test_text_runs_and_spacing(self):
        payload = page(
            [
                {
                    "id": "t",
                    "type": "text",
                    "bbox": [100, 100, 1000, 80],
                    "runs": [{"text": "Plain "}, {"text": "bold", "bold": True, "color": "#F05A28"}],
                    "style": {"font_size_pt": 24, "line_spacing_multiple": 1.25, "wrap": False, "margin_pt": 3.6},
                }
            ]
        )
        prs = compile_ir(payload, self.tmp)
        (shape,) = shape_by_name_order(prs)
        runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
        self.assertEqual([r.text for r in runs], ["Plain ", "bold"])
        self.assertTrue(runs[1].font.bold)
        self.assertEqual(str(runs[1].font.color.rgb), "F05A28")
        body = shape._element.find(".//a:bodyPr", NS)
        self.assertEqual(body.get("wrap"), "none")
        # 3.6 pt margin = 45720 EMU.
        self.assertEqual(int(body.get("lIns")), 45720)
        self.assertIsNotNone(shape._element.find(".//a:lnSpc/a:spcPct", NS))

    def test_portrait_page_keeps_ratio(self):
        payload = page([], width=900, height=1600)
        prs = compile_ir(payload, self.tmp)
        self.assertAlmostEqual(Emu(prs.slide_height).inches, 7.5, places=2)
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 900 / 1600, places=3)

    def test_features_example_compiles_and_passes_inspection(self):
        sys.path.insert(0, str(ROOT / "scripts"))
        import inspect_pptx

        payload = json.loads((ROOT / "examples" / "features" / "page_ir.json").read_text(encoding="utf-8"))
        compile_ir(payload, self.tmp)
        report = inspect_pptx.inspect(self.tmp / "out.pptx")
        self.assertTrue(report["pass"], report)


def xfrm(shape):
    return shape._element.find(".//a:xfrm", NS)


def page(objects, width=1600, height=900):
    return {
        "schema_version": "1.1",
        "pages": [{"id": "slide-1", "width_px": width, "height_px": height, "background": "#FFFFFF", "objects": objects}],
    }


if __name__ == "__main__":
    unittest.main()
