#!/usr/bin/env python3
"""Inspect PPTX geometry and flag full-slide raster shortcuts."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import numpy as np
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


FULL_SLIDE_COVERAGE = 0.95
COVERAGE_GRID = 400


def picture_union_coverage(boxes: list[tuple[int, int, int, int]], slide_w: int, slide_h: int) -> float:
    """Fraction of the slide covered by the union of picture boxes (catches tiled screenshots)."""
    if not boxes:
        return 0.0
    grid_w = COVERAGE_GRID
    grid_h = max(1, round(COVERAGE_GRID * slide_h / slide_w))
    mask = np.zeros((grid_h, grid_w), dtype=bool)
    for x, y, w, h in boxes:
        x0 = max(0, int(np.floor(x / slide_w * grid_w)))
        y0 = max(0, int(np.floor(y / slide_h * grid_h)))
        x1 = min(grid_w, int(np.ceil((x + w) / slide_w * grid_w)))
        y1 = min(grid_h, int(np.ceil((y + h) / slide_h * grid_h)))
        if x1 > x0 and y1 > y0:
            mask[y0:y1, x0:x1] = True
    return float(mask.mean())


def inspect(path: Path) -> dict:
    prs = Presentation(path)
    slide_w = int(prs.slide_width)
    slide_h = int(prs.slide_height)
    slides = []
    total_oob = 0
    total_full_slide_images = 0
    raster_shortcut_slides = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        text_shapes = 0
        picture_shapes = 0
        editable_chars = 0
        out_of_bounds = []
        full_slide_images = []
        picture_boxes = []
        for shape_index, shape in enumerate(slide.shapes, start=1):
            x, y, w, h = int(shape.left), int(shape.top), int(shape.width), int(shape.height)
            if x < 0 or y < 0 or x + w > slide_w or y + h > slide_h:
                out_of_bounds.append(shape_index)
            if getattr(shape, "has_text_frame", False):
                text = shape.text or ""
                if text.strip():
                    text_shapes += 1
                    editable_chars += len(text)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_shapes += 1
                picture_boxes.append((x, y, w, h))
                coverage = (w * h) / max(slide_w * slide_h, 1)
                if coverage >= FULL_SLIDE_COVERAGE:
                    full_slide_images.append(shape_index)
        union_coverage = picture_union_coverage(picture_boxes, slide_w, slide_h)
        if full_slide_images or union_coverage >= FULL_SLIDE_COVERAGE:
            raster_shortcut_slides.append(slide_index)
        total_oob += len(out_of_bounds)
        total_full_slide_images += len(full_slide_images)
        slides.append(
            {
                "slide": slide_index,
                "shape_count": len(slide.shapes),
                "text_shape_count": text_shapes,
                "editable_text_chars": editable_chars,
                "picture_count": picture_shapes,
                "picture_union_coverage": round(union_coverage, 4),
                "out_of_bounds_shape_indices": out_of_bounds,
                "full_slide_picture_indices": full_slide_images,
            }
        )
    return {
        "file": str(path),
        "slide_count": len(prs.slides),
        "slide_width_emu": slide_w,
        "slide_height_emu": slide_h,
        "out_of_bounds_count": total_oob,
        "full_slide_picture_count": total_full_slide_images,
        "raster_shortcut_slides": raster_shortcut_slides,
        "slides": slides,
        "pass": total_oob == 0 and not raster_shortcut_slides,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx")
    parser.add_argument("--json", action="store_true", help="accepted for compatibility; output is always JSON")
    args = parser.parse_args()
    report = inspect(Path(args.pptx))
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["pass"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
