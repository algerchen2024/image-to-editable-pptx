#!/usr/bin/env python3
"""Compile PageIR into an editable PPTX with python-pptx (no Node.js needed).

    compile_page_ir.py page_ir.json output.pptx

This is the default compiler: python-pptx is available in restricted sandboxes
(for example ChatGPT code execution) where Node packages cannot be installed.
It produces the same objects as compile_page_ir.js.
"""

from __future__ import annotations

import argparse
import copy
import json
import math
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pageir_common import DEFAULT_FONTS, latin_font_for, normalize_hex, slide_size_inches  # noqa: E402


EMU_PER_INCH = 914400
SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "round_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "chevron": MSO_SHAPE.CHEVRON,
}
DASH = {"solid": "solid", "dash": "dash", "dot": "sysDot", "dash_dot": "dashDot"}
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
VALIGN = {"top": MSO_ANCHOR.TOP, "mid": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def rgb(value, fallback: str) -> RGBColor:
    return RGBColor.from_string(normalize_hex(value if value else fallback)[1:])


def _sub(parent, tag: str, **attrs):
    el = etree.SubElement(parent, qn(tag))
    for key, val in attrs.items():
        el.set(key, str(val))
    return el


def set_line(shape_line, color, width_pt: float, dash: str | None) -> None:
    shape_line.color.rgb = rgb(color, "#111111")
    shape_line.width = Pt(width_pt)
    ln = shape_line._get_or_add_ln()
    ln.set("cap", "rnd")
    for old in ln.findall(qn("a:prstDash")):
        ln.remove(old)
    prst = etree.Element(qn("a:prstDash"))
    prst.set("val", DASH.get(dash or "solid", "solid"))
    # prstDash must follow the fill element inside a:ln.
    fill = ln.find(qn("a:solidFill"))
    if fill is not None:
        fill.addnext(prst)
    else:
        ln.insert(0, prst)


def drop_theme_style(shape) -> None:
    """Remove python-pptx's theme style reference (it adds a drop shadow and theme
    colors); every fill, line and font is set explicitly instead."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def set_rotation(shape, deg) -> None:
    if deg:
        shape.rotation = float(deg)


def apply_run_style(run, style: dict, base: dict, lang: str, platform: str, any_mode: bool) -> None:
    merged = {**base, **{k: v for k, v in style.items() if v is not None}}
    font = run.font
    font.size = Pt(float(merged.get("font_size_pt", 18)))
    font.bold = bool(merged.get("bold", False))
    font.italic = bool(merged.get("italic", False))
    font.underline = bool(merged.get("underline", False))
    font.color.rgb = rgb(merged.get("color"), "#111111")
    face = merged.get("font_face") or base.get("font_face") or DEFAULT_FONTS.get(platform, DEFAULT_FONTS["any"])["cjk"]
    latin = latin_font_for(merged, platform, str(face))
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", lang)
    if any_mode:
        rpr.set("altLang", "en-US")
    if merged.get("char_spacing_pt") is not None:
        rpr.set("spc", str(int(round(float(merged["char_spacing_pt"]) * 100))))
    for tag in ("a:latin", "a:ea", "a:cs"):
        for old in rpr.findall(qn(tag)):
            rpr.remove(old)
    _sub(rpr, "a:ea", typeface=face)
    _sub(rpr, "a:latin", typeface=latin)
    _sub(rpr, "a:cs", typeface=latin)


def add_text(slide, obj: dict, box, lang: str, platform: str, any_mode: bool) -> None:
    style = obj.get("style") or {}
    shape = slide.shapes.add_textbox(*box)
    tf = shape.text_frame
    tf.word_wrap = style.get("wrap", True) is not False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    margin = Pt(float(style.get("margin_pt", 0)))
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = margin
    tf.vertical_anchor = VALIGN.get(style.get("valign", "top"), MSO_ANCHOR.TOP)
    base = {
        "font_face": style.get("font_face") or DEFAULT_FONTS.get(platform, DEFAULT_FONTS["any"])["cjk"],
        "latin_font_face": style.get("latin_font_face"),
        "font_size_pt": style.get("font_size_pt", 18),
        "bold": style.get("bold", False),
        "italic": style.get("italic", False),
        "underline": style.get("underline", False),
        "color": style.get("color") or "#111111",
        "char_spacing_pt": style.get("char_spacing_pt"),
    }
    # Split runs into paragraphs at newlines.
    runs = obj.get("runs") if isinstance(obj.get("runs"), list) else [{"text": obj.get("text") or ""}]
    paragraphs: list[list[tuple[str, dict]]] = [[]]
    for run in runs:
        parts = str(run.get("text", "")).split("\n")
        for i, part in enumerate(parts):
            if i:
                paragraphs.append([])
            if part:
                paragraphs[-1].append((part, run))
    for p_index, items in enumerate(paragraphs):
        para = tf.paragraphs[0] if p_index == 0 else tf.add_paragraph()
        para.alignment = ALIGN.get(style.get("align", "left"), PP_ALIGN.LEFT)
        para.line_spacing = float(style.get("line_spacing_multiple", 1.0))
        para.space_before = Pt(0)
        para.space_after = Pt(0)
        if not items:
            # Keep empty lines at the right size.
            end = para._p.get_or_add_endParaRPr()
            end.set("sz", str(int(float(base["font_size_pt"]) * 100)))
        for text, run_style in items:
            run = para.add_run()
            run.text = text
            apply_run_style(
                run, {k: v for k, v in run_style.items() if k != "text"}, base, lang, platform, any_mode
            )
    set_rotation(shape, obj.get("rotation_deg"))


def add_shape(slide, obj: dict, box, scale: float) -> None:
    style = obj.get("style") or {}
    shape = slide.shapes.add_shape(SHAPES[obj["type"]], *box)
    if style.get("fill", "#FFFFFF") is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(style.get("fill"), "#FFFFFF")
        transparency = float(style.get("fill_transparency", 0) or 0)
        if transparency:
            clr = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            _sub(clr, "a:alpha", val=int(round((100 - transparency) * 1000)))
    if style.get("line", "#111111") is None:
        shape.line.fill.background()
    else:
        set_line(shape.line, style.get("line"), float(style.get("line_width_pt", 1)), style.get("line_dash"))
    if obj["type"] == "round_rect" and style.get("corner_radius_px") is not None:
        radius_emu = float(style["corner_radius_px"]) * scale
        shortest = max(1, min(box[2], box[3]))
        shape.adjustments[0] = max(0.0, min(0.5, radius_emu / shortest))
    drop_theme_style(shape)
    set_rotation(shape, obj.get("rotation_deg"))


def _positive_pair(value) -> tuple[float, float] | None:
    if (
        isinstance(value, (list, tuple))
        and len(value) == 2
        and all(isinstance(v, (int, float)) and not isinstance(v, bool) and float(v) > 0 for v in value)
    ):
        return float(value[0]), float(value[1])
    return None


def arrow_head_placement(x1: float, y1: float, x2: float, y2: float, head_w: float, head_l: float) -> dict | None:
    """Shaft stops at the head base; the triangle tip sits on (x2, y2).

    `head_w` is the base and `head_l` the altitude, both in source pixels.
    Rotation is clockwise degrees from the preset triangle's upward tip.
    """
    dx, dy = x2 - x1, y2 - y1
    dist = math.hypot(dx, dy)
    if dist < 1:
        return None
    ux, uy = dx / dist, dy / dist
    length = min(float(head_l), dist * 0.85)
    base = (x2 - ux * length, y2 - uy * length)
    center = (x2 - ux * length / 2.0, y2 - uy * length / 2.0)
    return {
        "shaft": (x1, y1, base[0], base[1]),
        "center": center,
        "width": float(head_w),
        "height": length,
        "rotation": math.degrees(math.atan2(ux, -uy)) % 360.0,
    }


def _add_arrow_head(slide, place: dict, color, sx: float, sy: float) -> None:
    """Native solid isosceles triangle. Local width is the base, height the altitude."""
    cx, cy = place["center"]
    scale = (sx + sy) / 2.0
    width = place["width"] * scale
    height = max(place["height"] * scale, 1.0)
    left = cx * sx - width / 2.0
    top = cy * sy - height / 2.0
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        Emu(int(round(left))),
        Emu(int(round(top))),
        Emu(int(round(width))),
        Emu(int(round(max(height, 1.0)))),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color, "#111111")
    shape.line.fill.background()
    drop_theme_style(shape)
    set_rotation(shape, place["rotation"])


def _add_open_arrow_wings(slide, x1, y1, x2, y2, width_px, length_px, color, width_pt, dash, sx, sy):
    dx, dy = x2 - x1, y2 - y1
    dist = math.hypot(dx, dy)
    if dist < 1:
        return
    ux, uy = dx / dist, dy / dist
    px, py = -uy, ux
    length = min(float(length_px), dist * 0.85)
    bx, by = x2 - ux * length, y2 - uy * length
    for ex, ey in ((bx + px * width_px / 2.0, by + py * width_px / 2.0),
                   (bx - px * width_px / 2.0, by - py * width_px / 2.0)):
        wing = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Emu(int(round(ex * sx))), Emu(int(round(ey * sy))),
            Emu(int(round(x2 * sx))), Emu(int(round(y2 * sy))),
        )
        drop_theme_style(wing)
        set_line(wing.line, color, width_pt, dash)


def add_line(slide, obj: dict, sx: float, sy: float) -> None:
    style = obj.get("style") or {}
    x1, y1, x2, y2 = (float(v) for v in obj["points"])
    head = _positive_pair(style.get("head_size_px"))
    place = None
    if head:
        # A start-only arrow's tip is the first point; otherwise the tip is the end point.
        start_only = style.get("start_arrow") not in (None, "none") and style.get("end_arrow") in (None, "none")
        if start_only:
            place = arrow_head_placement(x2, y2, x1, y1, head[0], head[1])
        else:
            place = arrow_head_placement(x1, y1, x2, y2, head[0], head[1])
    if place is not None and (
        style.get("end_arrow") == "open" or style.get("start_arrow") == "open"
    ):
        # Open heads keep the shaft all the way to the tip and add two native
        # line segments for the chevron wings.
        conn = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Emu(int(x1 * sx)), Emu(int(y1 * sy)), Emu(int(x2 * sx)), Emu(int(y2 * sy))
        )
        drop_theme_style(conn)
        set_line(conn.line, style.get("color"), float(style.get("width_pt", 1)), style.get("dash"))
        if style.get("end_arrow") == "open":
            _add_open_arrow_wings(
                slide, x1, y1, x2, y2, head[0], head[1], style.get("color"), float(style.get("width_pt", 1)),
                style.get("dash"), sx, sy
            )
        if style.get("start_arrow") == "open":
            _add_open_arrow_wings(
                slide, x2, y2, x1, y1, head[0], head[1], style.get("color"), float(style.get("width_pt", 1)),
                style.get("dash"), sx, sy
            )
        return
    if place is not None:
        x1, y1, x2, y2 = place["shaft"]
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Emu(int(x1 * sx)), Emu(int(y1 * sy)), Emu(int(x2 * sx)), Emu(int(y2 * sy))
    )
    drop_theme_style(conn)
    set_line(conn.line, style.get("color"), float(style.get("width_pt", 1)), style.get("dash"))
    if place is not None:
        # Custom triangle replaces DrawingML headEnd/tailEnd. Tip is the line's end point.
        _add_arrow_head(slide, place, style.get("color"), sx, sy)
        return
    ln = conn.line._get_or_add_ln()
    for key, tag in (("start_arrow", "a:headEnd"), ("end_arrow", "a:tailEnd")):
        if style.get(key) in {"triangle", "open"}:
            attrs = {"type": "arrow" if style.get(key) == "open" else "triangle"}
            if style.get("head_w") in {"sm", "med", "lg"}:
                attrs["w"] = style["head_w"]
            if style.get("head_len") in {"sm", "med", "lg"}:
                attrs["len"] = style["head_len"]
            _sub(ln, tag, **attrs)


def _patch_any_theme(output: Path) -> None:
    """Use opening-machine defaults for CJK and Arial for theme Latin text."""
    patched = output.with_suffix(".theme-patched.pptx")
    with zipfile.ZipFile(output, "r") as zin, zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/theme/theme1.xml":
                root = etree.fromstring(data)
                for tag in ("a:majorFont", "a:minorFont"):
                    block = root.find(f".//{qn(tag)}")
                    if block is None:
                        continue
                    ea = block.find(qn("a:ea"))
                    if ea is not None:
                        ea.set("typeface", "")
                    for slot in ("a:latin", "a:cs"):
                        node = block.find(qn(slot))
                        if node is not None:
                            node.set("typeface", "Arial")
                    for font in list(block):
                        if font.tag == qn("a:font") and font.get("script") in {"Hans", "Hant", "Jpan", "Hang"}:
                            block.remove(font)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    patched.replace(output)


def add_image(slide, obj: dict, box, base_dir: Path) -> None:
    path = Path(obj["path"])
    if not path.is_absolute():
        path = (base_dir / path).resolve()
    pic = slide.shapes.add_picture(str(path), *box)
    set_rotation(pic, obj.get("rotation_deg"))


def compile_payload(payload: dict, base_dir: Path, output: Path) -> dict:
    pages = payload.get("pages") or []
    if not pages:
        raise ValueError("PageIR contains no pages")
    first = pages[0]
    aspect = float(first["width_px"]) / float(first["height_px"])
    width_in, height_in = slide_size_inches(float(first["width_px"]), float(first["height_px"]))
    lang = payload.get("lang") or "zh-CN"
    target = payload.get("target")
    platform = str(("mac" if target is None else (target or {}).get("platform") or "any")).lower()
    any_mode = platform == "any"

    prs = Presentation()
    prs.slide_width = Emu(int(round(width_in * EMU_PER_INCH)))
    prs.slide_height = Emu(int(round(height_in * EMU_PER_INCH)))
    prs.core_properties.title = payload.get("title") or "Editable PowerPoint reconstruction"
    prs.core_properties.author = "image-to-editable-pptx"
    blank = prs.slide_layouts[6]
    for page in pages:
        if abs(float(page["width_px"]) / float(page["height_px"]) - aspect) / aspect > 0.001:
            raise ValueError("All pages must share the same aspect ratio")
        slide = prs.slides.add_slide(blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(page.get("background"), "#FFFFFF")
        sx = prs.slide_width / float(page["width_px"])
        sy = prs.slide_height / float(page["height_px"])
        page_lang = page.get("lang") or lang
        # Stable sort: equal z keeps PageIR order.
        for obj in sorted(page.get("objects") or [], key=lambda o: float(o.get("z", 0) or 0)):
            if obj["type"] == "line":
                add_line(slide, obj, sx, sy)
                continue
            x, y, w, h = (float(v) for v in obj["bbox"])
            box = (Emu(int(round(x * sx))), Emu(int(round(y * sy))), Emu(int(round(w * sx))), Emu(int(round(h * sy))))
            if obj["type"] == "text":
                add_text(slide, obj, box, page_lang, platform, any_mode)
            elif obj["type"] == "image":
                add_image(slide, obj, box, base_dir)
            else:
                add_shape(slide, obj, box, min(sx, sy))
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    if any_mode:
        _patch_any_theme(output)
    return {"output": str(output), "slides": len(pages), "width_in": width_in, "height_in": height_in}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("page_ir")
    parser.add_argument("output")
    args = parser.parse_args()
    ir_path = Path(args.page_ir).resolve()
    payload = json.loads(ir_path.read_text(encoding="utf-8"))
    try:
        result = compile_payload(copy.deepcopy(payload), ir_path.parent, Path(args.output).resolve())
    except (ValueError, KeyError) as exc:
        print(str(exc), file=sys.stderr)
        return 2
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
